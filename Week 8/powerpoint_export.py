
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt


def create_powerpoint(insights, executive_summary, recommendations):
    """Create a seven-slide executive presentation and return its bytes."""

    if not isinstance(insights, (list, tuple)) or len(insights) != 5:
        raise ValueError("Exactly five insights are required.")

    if not isinstance(executive_summary, str) or not executive_summary.strip():
        raise ValueError("An executive summary is required.")

    if (
        not isinstance(recommendations, (list, tuple))
        or not recommendations
        or any(not isinstance(item, str) or not item.strip()
               for item in recommendations)
    ):
        raise ValueError("At least one valid recommendation is required.")

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    dark_blue = RGBColor(35, 48, 68)
    grey = RGBColor(100, 110, 125)

    def add_text(
        slide, text, left, top, width, height,
        font_size=24, bold=False, color=dark_blue
    ):
        box = slide.shapes.add_textbox(
            Inches(left), Inches(top), Inches(width), Inches(height)
        )

        frame = box.text_frame
        frame.word_wrap = True
        frame.margin_left = Inches(0.12)
        frame.margin_right = Inches(0.12)
        frame.margin_top = Inches(0.08)
        frame.margin_bottom = Inches(0.08)

        paragraph = frame.paragraphs[0]
        paragraph.text = str(text)
        paragraph.font.size = Pt(font_size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color

        return box

    def add_footer(slide, slide_number):
        add_text(
            slide,
            "CIA Executive Insights | Week 8 Day 3",
            0.7, 7.05, 10.5, 0.25,
            font_size=10, color=grey,
        )
        add_text(
            slide,
            f"{slide_number} / 7",
            11.5, 7.05, 1.0, 0.25,
            font_size=10, color=grey,
        )

    # Slide 1: Title and executive summary
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(
        slide, "Executive Business Insights",
        0.7, 0.7, 12, 0.8, font_size=30, bold=True,
    )
    add_text(
        slide, "EDA Report | Week 8 Day 3",
        0.8, 1.7, 11.5, 0.5, font_size=18,
    )
    add_text(
        slide, executive_summary,
        0.8, 2.5, 11.5, 4.2, font_size=17,
    )
    add_footer(slide, 1)

    # Slides 2–6: Five insights
    for index, insight in enumerate(insights, start=1):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        add_text(
            slide, f"Insight {index}",
            0.7, 0.7, 11.8, 0.8, font_size=30, bold=True,
        )
        add_text(
            slide, insight,
            0.9, 2.0, 11.2, 3.8, font_size=23,
        )
        add_footer(slide, index + 1)

    # Slide 7: Business recommendations
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_text(
        slide, "Business Recommendations",
        0.7, 0.7, 12, 0.8, font_size=30, bold=True,
    )

    recommendation_text = "\n".join(
        f"{index}. {item}"
        for index, item in enumerate(recommendations, start=1)
    )
    add_text(
        slide, recommendation_text,
        0.9, 1.8, 11.5, 4.8, font_size=19,
    )
    add_footer(slide, 7)

    output = BytesIO()
    prs.save(output)
    output.seek(0)
    return output.getvalue()
